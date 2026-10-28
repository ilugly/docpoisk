"""Тесты чтения форматов и нарезки на чанки (без нейросетей). Запуск: pytest backend/tests"""
import io

import pytest

from app.chunking import chunk_pages, normalize, split_text
from app.parsing import UnsupportedFormat, extract_pages, unsupported_reason


def test_split_respects_size_and_overlap():
    text = " ".join(f"Предложение номер {i}." for i in range(300))
    chunks = split_text(text, 200, 50)
    assert len(chunks) > 5
    assert all(len(c) <= 200 for c in chunks)
    # соседние чанки перекрываются
    assert any(w in chunks[1] for w in chunks[0].split()[-3:])


def test_short_text_is_single_chunk():
    assert split_text("Короткий текст", 1000, 200) == ["Короткий текст"]


def test_normalize_joins_hyphenated_words():
    assert normalize("дого-\nвор  подряда\r\n\r\n\r\nпункт") == "договор подряда\n\nпункт"


def test_chunk_pages_keeps_page_numbers():
    chunks = chunk_pages([(1, "а " * 600), (2, "б " * 10)], 500, 50)
    assert chunks[0].page == 1 and chunks[-1].page == 2
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_docx_with_table():
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("Договор подряда № 15")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Подсудность"
    t.cell(0, 1).text = "Арбитражный суд г. Москвы"
    buf = io.BytesIO()
    d.save(buf)
    [(_, text)] = extract_pages("dogovor.docx", buf.getvalue())
    assert "Договор подряда № 15" in text
    assert "Подсудность | Арбитражный суд г. Москвы" in text


def test_xlsx_sheets():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Этапы"
    wb.active.append(["Этап 1", "Проектирование"])
    buf = io.BytesIO()
    wb.save(buf)
    [(page, text)] = extract_pages("plan.xlsx", buf.getvalue())
    assert page == 1 and "Лист: Этапы" in text and "Этап 1 | Проектирование" in text


def test_txt_cp1251():
    [(_, text)] = extract_pages("old.txt", "Срок оплаты — 30 дней".encode("cp1251"))
    assert text == "Срок оплаты — 30 дней"


def test_html_strips_scripts():
    [(_, text)] = extract_pages("a.html", b"<html><script>x()</script><p>Hello</p></html>")
    assert "Hello" in text and "x()" not in text


def test_unsupported():
    assert "docx" in unsupported_reason("old.doc")
    assert unsupported_reason("a.pdf") is None
    with pytest.raises(UnsupportedFormat):
        extract_pages("image.png", b"")
