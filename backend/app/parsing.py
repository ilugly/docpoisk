"""Извлечение текста из документов разных форматов.

Каждый парсер возвращает список «страниц»: (номер_страницы | None, текст).
Номер страницы нужен, чтобы в ответе можно было сослаться на источник.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Callable

Page = tuple[int | None, str]


class UnsupportedFormat(Exception):
    pass


def _pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # битая страница не должна ронять весь документ
            text = ""
        pages.append((i, text))
    return pages


def _docx(data: bytes) -> list[Page]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    # Идём по телу документа по порядку, чтобы таблицы оставались на своих местах
    for block in doc.element.body.iterchildren():
        tag = block.tag.split("}")[-1]
        if tag == "p":
            text = Paragraph(block, doc).text.strip()
            if text:
                parts.append(text)
        elif tag == "tbl":
            for row in Table(block, doc).rows:
                cells = []
                for cell in row.cells:
                    t = cell.text.strip()
                    if t and (not cells or cells[-1] != t):  # объединённые ячейки повторяются
                        cells.append(t)
                if cells:
                    parts.append(" | ".join(cells))
    return [(None, "\n".join(parts))]


def _xlsx(data: bytes) -> list[Page]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    for idx, ws in enumerate(wb.worksheets, start=1):
        lines = [f"Лист: {ws.title}"]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        pages.append((idx, "\n".join(lines)))
    return pages


def _pptx(data: bytes) -> list[Page]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        pages.append((i, "\n".join(t for t in texts if t.strip())))
    return pages


def _odt(data: bytes) -> list[Page]:
    from odf import teletype
    from odf.opendocument import load
    from odf.text import H, P

    doc = load(io.BytesIO(data))
    parts = [teletype.extractText(el) for el in doc.getElementsByType(H) + doc.getElementsByType(P)]
    return [(None, "\n".join(p for p in parts if p.strip()))]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "koi8-r"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _txt(data: bytes) -> list[Page]:
    return [(None, _decode(data))]


def _csv(data: bytes) -> list[Page]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [" | ".join(c.strip() for c in r if c.strip()) for r in csv.reader(io.StringIO(text), dialect)]
    return [(None, "\n".join(r for r in rows if r))]


def _json(data: bytes) -> list[Page]:
    try:
        return [(None, json.dumps(json.loads(_decode(data)), ensure_ascii=False, indent=1))]
    except ValueError:
        return _txt(data)


def _html(data: bytes) -> list[Page]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return [(None, soup.get_text("\n"))]


def _rtf(data: bytes) -> list[Page]:
    from striprtf.striprtf import rtf_to_text

    return [(None, rtf_to_text(_decode(data), errors="ignore"))]


PARSERS: dict[str, Callable[[bytes], list[Page]]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".pptx": _pptx,
    ".odt": _odt,
    ".txt": _txt,
    ".md": _txt,
    ".log": _txt,
    ".xml": _txt,
    ".csv": _csv,
    ".json": _json,
    ".html": _html,
    ".htm": _html,
    ".rtf": _rtf,
}

SUPPORTED = sorted(PARSERS)


def unsupported_reason(filename: str) -> str | None:
    """None, если формат поддерживается, иначе — понятное объяснение."""
    ext = Path(filename).suffix.lower()
    if ext in PARSERS:
        return None
    hint = ""
    if ext in {".doc", ".xls", ".ppt"}:
        hint = " Старые форматы Office: пересохраните файл как .docx/.xlsx/.pptx."
    return f"Формат {ext or '(без расширения)'} не поддерживается.{hint}"


def extract_pages(filename: str, data: bytes) -> list[Page]:
    reason = unsupported_reason(filename)
    if reason:
        raise UnsupportedFormat(reason)
    parser = PARSERS[Path(filename).suffix.lower()]
    return [(n, t) for n, t in parser(data) if t and t.strip()]
